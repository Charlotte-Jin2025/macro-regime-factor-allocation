"""Build the single reviewed report from the current analysis and manuscript.

All paths are relative to this delivery directory. No earlier Word report,
historical output directory, or locally installed office application is used.
"""
from __future__ import annotations

import csv
import base64
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import json
import math
from pathlib import Path
import re
from statistics import mean
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CURRENT = "Release_corrected_current_vintage"
PRIMARY = "LS5_Variance36_Cap40_Shrink50"
FACTORS = ["MKT_RF", "SMB", "HML", "RMW", "CMA", "MOM"]
STRATEGIES = ["Equal_Weight", "Inverse_Volatility", "Historical_Mean", "Macro", "Macro_M1_YoY", "Macro_M1_ExpandingYoY"]
LABEL = {
    "Equal_Weight": "Equal weight", "Inverse_Volatility": "Inverse volatility",
    "Historical_Mean": "Historical mean", "Macro": "Macro",
    "Macro_M1_YoY": "Fixed state", "Macro_M1_ExpandingYoY": "Expanding state",
    "Macro_minus_Historical_Mean": "Macro minus historical mean",
    "Macro_M1_YoY_minus_Macro": "Fixed state minus Macro",
    "Macro_M1_ExpandingYoY_minus_Macro": "Expanding state minus Macro",
}
MAPPING_LABEL = {
    "Raw6_Sensitivity": "Six factors / raw",
    "Raw6_VolCeiling5": "Six factors / raw / 5% ceiling",
    "Risk6_Cap40_Shrink50": "Six factors / variance-scaled",
    "Risk6_Cap40_Shrink50_VolCeiling5": "Six factors / variance / 5% ceiling",
    "Rank6_Cap40_Shrink50": "Six factors / ranks",
    "Rank6_Cap40_Shrink50_VolCeiling5": "Six factors / ranks / 5% ceiling",
    "LS5_PositiveRaw": "Five factors / raw",
    PRIMARY: "Five factors / variance (primary)",
    "FixedMarket1of6_LS5_PositiveRaw": "Fixed market 1/6 / raw",
    "FixedMarket1of6_LS5_Variance36_Cap40_Shrink50": "Fixed market 1/6 / variance",
}


def load_csv(base: Path, relative: str) -> list[dict]:
    with (base / relative).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def select(rows: list[dict], **filters) -> list[dict]:
    return [row for row in rows if all(str(row.get(k, "")) == str(v) for k, v in filters.items())]


def unique(rows: list[dict], **filters) -> dict:
    found = select(rows, **filters)
    if len(found) != 1:
        raise ValueError(f"Expected one result for {filters}; found {len(found)}")
    return found[0]


def number(x, digits=3, scale=1.0) -> str:
    value = float(x) * scale
    if abs(value) < 0.5 * 10 ** (-digits):
        value = 0.0
    return f"{value:.{digits}f}".replace("-", "−")


def format_r2(x) -> str:
    value = float(x) * 100
    if 0 < abs(value) < 0.005:
        return "+<0.01" if value > 0 else "−<0.01"
    return number(value, 2)


def format_rms_bp(x) -> str:
    value=float(x)
    return "<0.01" if 0 < value < .01 else number(value,2)


def model_r2(rows: list[dict], column: str, benchmark="Historical_Mean") -> float:
    errors = [(float(r["Actual"]) - float(r[column])) ** 2 for r in rows]
    base = [(float(r["Actual"]) - float(r[benchmark])) ** 2 for r in rows]
    return 1 - sum(errors) / sum(base)


def data_context(base: Path):
    datasets = {
        "metrics": load_csv(base, "results/us/metrics.csv"),
        "intervals": load_csv(base, "results/us/studentized_intervals.csv"),
        "mde": load_csv(base, "results/us/mde.csv"),
        "mean_tests": load_csv(base, "results/us/mean_difference_hac.csv"),
        "forecasts": load_csv(base, "results/us/forecasts.csv"),
        "h1": load_csv(base, "results/us/direct_h1_forecast_loss.csv"),
        "h2": load_csv(base, "results/us/direct_h2_forecast_loss.csv"),
        "states": load_csv(base, "results/states/state_frequencies.csv"),
        "calendar": load_csv(base, "results/states/model_calendar.csv"),
        "mom": load_csv(base, "results/diagnostics/mom_internal_cost_sensitivity.csv"),
        "exposure": load_csv(base, "results/us/exposure_timing_comparison.csv"),
        "spanning": load_csv(base, "results/us/spanning_hac.csv"),
        "event": load_csv(base, "results/diagnostics/cumulative_portfolio_difference.csv"),
        "intl": load_csv(base, "results/international/model_shrinkage_diagnostics.csv"),
        "complexity": load_csv(base, "results/diagnostics/initial_effective_df.csv"),
        "mapping_comparisons": load_csv(base, "results/selection/mapping_comparisons.csv"),
        "reality_check": load_csv(base, "results/selection/reality_check.csv"),
        "annual": load_csv(base, "results/selection/annual_returns.csv"),
        "covid_mappings": load_csv(base, "results/interpretation/event_2020_mapping_comparison.csv"),
        "annual_increment": load_csv(base, "results/robustness/annual_incremental_returns.csv"),
        "annual_concentration": load_csv(base, "results/robustness/annual_concentration_summary.csv"),
        "ip_metrics": load_csv(base, "results/robustness/ip_design_metrics.csv"),
        "ip_mean_tests": load_csv(base, "results/robustness/ip_design_mean_hac.csv"),
        "ip_intervals": load_csv(base, "results/robustness/ip_design_sharpe_intervals.csv"),
        "ip_decomposition": load_csv(base, "results/robustness/ip_sequential_decomposition.csv"),
        "training_slopes": load_csv(base, "results/robustness/training_slopes.csv"),
        "rmw_state_means": load_csv(base, "results/robustness/rmw_state_means.csv"),
        "mapping_turnover": load_csv(base, "results/robustness/mapping_turnover.csv"),
        "cost_crossings": load_csv(base, "results/robustness/cost_crossings.csv"),
        "month_influence": load_csv(base, "results/robustness/month_influence_extremes.csv"),
    }
    selection_summary=json.loads((base / "results/selection/summary.json").read_text(encoding="utf-8"))
    datasets["selection_summary"]=selection_summary
    common = dict(Design=CURRENT, Cohort="Extended_175", Allocation=PRIMARY, Cost_bps="10", Sample="Full")
    perf = {s: unique(datasets["metrics"], Strategy=s, **common) for s in STRATEGIES}
    comparison = "Macro_minus_Historical_Mean"
    ci = unique(datasets["intervals"], Comparison=comparison, Block_Length="6", HAC_Lags="3", **common)
    mde = unique(datasets["mde"], Comparison=comparison, HAC_Lags="3", **common)
    mt = unique(datasets["mean_tests"], Comparison=comparison, HAC_Lags="3", **common)
    c = select(datasets["calendar"], Design=CURRENT)
    train = [r for r in c if "1990-03-31" <= r["Signal_Date"] <= "2011-11-30"]
    evaluation = [r for r in c if "2019-03-31" <= r["Signal_Date"] <= "2026-06-30"]
    observed = [float(r["IP_YoY_lag"]) for r in evaluation if r["IP_YoY_lag"]]
    first, last = train[0], evaluation[-1]
    forecasts = select(datasets["forecasts"], Design=CURRENT, Factor="RMW")
    observed_forecasts = [r for r in forecasts if r["H2_YoY_OLS_Eligible"] == "True"]
    rmw = unique(datasets["h2"], Design=CURRENT, Cohort="Extended_175", State_Model="Macro_M1_YoY", Factor="RMW", Estimator="OLS", Method="Clark_West", HAC_Lags="3")
    rmw_loss = unique(datasets["h2"], Design=CURRENT, Cohort="Extended_175", State_Model="Macro_M1_YoY", Factor="RMW", Estimator="OLS", Method="DM_Loss_HAC", HAC_Lags="3")
    def rmw_test(cohort, state):
        return unique(datasets["h2"], Design=CURRENT, Cohort=cohort, State_Model=state, Factor="RMW", Estimator="OLS", Method="Clark_West", HAC_Lags="3")
    rmw_exp=rmw_test("Extended_175", "Macro_M1_ExpandingYoY")
    rmw_88=rmw_test("Original_88", "Macro_M1_YoY")
    rmw_exp_88=rmw_test("Original_88", "Macro_M1_ExpandingYoY")
    h1_min={cohort:min(float(row["Holm_p_Six_Factors"]) for row in select(datasets["h1"], Design=CURRENT, Cohort=cohort, Estimator="OLS", Method="Clark_West", HAC_Lags="3")) for cohort in ["Extended_175","Original_88"]}
    rc=unique(datasets["reality_check"], Cohort="Extended_175", Family="Macro_10_Mappings", Block_Length="6")
    annual_2022=unique(datasets["annual"], Year="2022", Strategy="Macro")
    ip_common = {**common, "Design": "Partial_realtime_IP"}
    ip_ci = unique(datasets["intervals"], Comparison=comparison, Block_Length="6", HAC_Lags="3", **ip_common)
    mom = {s: unique(datasets["mom"], Design=CURRENT, Cohort="Extended_175", Allocation=PRIMARY, Sample="Full", Strategy=s, Allocation_Cost_bps="10", MOM_Internal_Cost_bps_per_month="65.0") for s in ["Macro", "Historical_Mean"]}
    exposure = unique(datasets["exposure"], Comparison=comparison, Design=CURRENT, Cohort="Extended_175", Allocation=PRIMARY, Sample="Full")
    spanning = unique(datasets["spanning"], Comparison=comparison, HAC_Lags="3", **common)
    no_event = unique(datasets["mean_tests"], Comparison=comparison, HAC_Lags="3", **{**common, "Sample":"Exclude_2020_Mar_Dec"})
    events = select(datasets["event"], Design=CURRENT, Cohort="Extended_175", Comparison=comparison, Cost_bps="10", Sample="Full")
    event_primary = [r for r in events if r["Allocation"] == PRIMARY and "2020-03-31" <= r["Return_Date"] <= "2020-12-31"]
    event_raw = [r for r in events if r["Allocation"] == "Raw6_Sensitivity" and "2020-03-31" <= r["Return_Date"] <= "2020-12-31"]
    assert len(event_primary) == len(event_raw) == 10
    values = {
        "macro_sr": number(perf["Macro"]["Sharpe"]), "hm_sr": number(perf["Historical_Mean"]["Sharpe"]),
        "delta_sr": number(ci["Sharpe_Difference"]), "ci_low": number(ci["CI_Low"]), "ci_high": number(ci["CI_High"]),
        "mde": number(mde["Approx_MDE_Annual_Sharpe_80pct"]),
        "macro_mean": number(perf["Macro"]["Annualized_Mean"], 2, 100), "macro_cagr": number(perf["Macro"]["Synthetic_CAGR"], 2, 100),
        "ew_mean": number(perf["Equal_Weight"]["Annualized_Mean"], 2, 100), "ew_sr": number(perf["Equal_Weight"]["Sharpe"]),
        "mean_gain": number(mt["Mean"], 3, 12), "mean_p": number(mt["p_Two_Sided"], 3), "mean_holm": number(mt["Holm_p_Within_Displayed_Comparison_Family"], 3),
        "growth_threshold": number(first["IP_YoY_lag_Fixed_Median"], 2), "inflation_threshold": number(first["CPI_YoY_lag_Fixed_Median"], 2),
        "ip_training_mean": number(mean(float(r["IP_YoY_lag"]) for r in train), 2), "ip_eval_mean": number(mean(observed), 2), "ip_eval_n": str(len(observed)),
        "growth_end_median": number(last["IP_YoY_lag_Expanding_Median"], 2),
        "rmw_reduced_r2": number(model_r2(observed_forecasts, "Macro_OLS_Matched_YoY_Labels"), 2, 100),
        "rmw_fixed_r2": number(model_r2(observed_forecasts, "Macro_M1_YoY_OLS"), 2, 100),
        "rmw_exp_r2": number(model_r2([r for r in forecasts if r["H2_ExpandingYoY_OLS_Eligible"] == "True"], "Macro_M1_ExpandingYoY_OLS"), 2, 100),
        "rmw_cw": number(rmw["Holm_p_Six_Factors"], 4), "rmw_loss": number(rmw_loss["Holm_p_Six_Factors"], 4),
        "rmw_exp_cw": number(rmw_exp["Holm_p_Six_Factors"], 4),
        "rmw_cw_88": number(rmw_88["Holm_p_Six_Factors"], 4), "rmw_exp_cw_88": number(rmw_exp_88["Holm_p_Six_Factors"], 3),
        "h1_min_175": number(h1_min["Extended_175"]), "h1_min_88": number(h1_min["Original_88"]),
        "fixed_expansion_validation": unique(datasets["states"], Design=CURRENT, Variant="FrozenYoY", Split="Validation", Regime="Expansion")["N"],
        "exp_expansion_validation": unique(datasets["states"], Design=CURRENT, Variant="ExpandingYoY", Split="Validation", Regime="Expansion")["N"],
        "ip_delta_sr": number(ip_ci["Sharpe_Difference"]), "ip_ci_low": number(ip_ci["CI_Low"]), "ip_ci_high": number(ip_ci["CI_High"]),
        "mom_macro_mean": number(mom["Macro"]["Annual_Mean_pp"], 2), "mom_hm_mean": number(mom["Historical_Mean"]["Annual_Mean_pp"], 2),
        "timing_difference": number(exposure["Annual_Dynamic_Timing_Return_pp"]), "static_difference": number(exposure["Annual_Static_Exposure_Return_pp"]),
        "alpha": number(spanning["Alpha_Annual_pp"]), "alpha_p": number(spanning["Alpha_p_Two_Sided"]),
        "no_covid_gain": number(no_event["Mean"], 3, 12),
        "covid_primary": number(sum(float(r["Monthly_Difference_pp"]) for r in event_primary), 2),
        "covid_raw": number(sum(float(r["Monthly_Difference_pp"]) for r in event_raw), 2),
        "macro_2022": number(annual_2022["Compounded_Return_pct"], 2),
        "macro_ex2022": number(selection_summary["non_2022_macro_compounded_return_pct"], 2),
        "rc_p": number(rc["Reality_Check_p"]),
        "rc_decision": "The family null is not rejected at 5%." if float(rc["Reality_Check_p"])>=.05 else "The recorded-family null is rejected at 5%, conditional on the specified paths and search family.",
    }
    cohort="Extended_175"
    annual=lambda year:unique(datasets["annual_increment"],Cohort=cohort,Year=str(year))
    concentration=lambda sample:unique(datasets["annual_concentration"],Cohort=cohort,Sample=sample)
    fixed="Partial_realtime_IP_Current_Penalties"
    ip_mean=lambda design:unique(datasets["ip_mean_tests"],Cohort=cohort,Design=design,HAC_Lags="3")
    ip_fixed_ci=unique(datasets["ip_intervals"],Cohort=cohort,Design=fixed)
    primary_turnover=unique(datasets["mapping_turnover"],Cohort=cohort,Mapping=PRIMARY)
    raw_turnover=unique(datasets["mapping_turnover"],Cohort=cohort,Mapping="LS5_PositiveRaw")
    crossing=lambda cost:unique(datasets["cost_crossings"],Cohort=cohort,Cost_Type=cost)
    august=unique(datasets["month_influence"],Cohort=cohort,Deleted_Return_Date="2020-08-31",Selection="Maximum_deleted_delta_Sharpe")
    retuning=unique(datasets["ip_decomposition"],Cohort=cohort,Step="Retuning_on_historical_IP_given_the_substitution")
    values.update({
        "annual_increment_total":number(concentration("Full")["Macro_minus_HM_Arithmetic_Sum_pp"],2),
        "annual_increment_2022":number(annual(2022)["Macro_minus_HM_Arithmetic_Sum_pp"],2),
        "annual_increment_2022_share":number(concentration("Only_2022")["Share_of_Full_Arithmetic_Increment_pct"],1),
        "annual_increment_2024":number(annual(2024)["Macro_minus_HM_Arithmetic_Sum_pp"],2),
        "ex2022_mean_gain":number(concentration("Exclude_2022")["Macro_minus_HM_Annual_Arithmetic_Mean_pp"],3),
        "hm_ex2022":number(concentration("Exclude_2022")["Historical_Mean_Compounded_Return_pct"],2),
        "hm_2022":number(annual(2022)["Historical_Mean_Compounded_Return_pct"],2),
        "ip_fixed_mean_gain":number(ip_mean(fixed)["Annual_Mean_Gain_pp"],3),
        "ip_retuned_mean_gain":number(ip_mean("Partial_realtime_IP")["Annual_Mean_Gain_pp"],3),
        "ip_retuning_increment":number(retuning["Annual_Arithmetic_Increment_Change_pp"],3),
        "ip_fixed_delta_sr":number(ip_fixed_ci["Sharpe_Difference"],3),
        "ip_fixed_ci_low":number(ip_fixed_ci["CI_Low"],3),
        "ip_fixed_ci_high":number(ip_fixed_ci["CI_High"],3),
        "allocation_crossing":number(crossing("Allocation_only")["First_Loss_of_Positive_Advantage_bps"],1),
        "mom_crossing":number(crossing("MOM_internal")["First_Loss_of_Positive_Advantage_bps"],1),
        "macro_turnover":number(primary_turnover["Macro_Mean_Monthly_One_Way_Turnover"],1,100),
        "hm_turnover":number(primary_turnover["HM_Mean_Monthly_One_Way_Turnover"],1,100),
        "raw5_delta_sr":number(raw_turnover["Sharpe_Difference"],3),
        "raw5_turnover":number(raw_turnover["Macro_Mean_Monthly_One_Way_Turnover"],1,100),
        "aug2020_deleted_delta_sr":number(august["Deleted_Sharpe_Difference"],3),
    })
    from .report_extension import load_extension
    load_extension(base, datasets, values)
    return datasets, perf, values


def table_data(kind, data, perf):
    """Return title, columns, rows, widths, note, and analysis sources."""
    if kind in {"cross_market_allocation", "international_sources", "international_allocation_uncertainty"}:
        from .report_extension import extension_table
        return extension_table(kind, data)
    if kind == "predictors":
        return ("Table 2 Seven predictors and their economic rationale", ["Predictor", "Observation used", "Economic channel"], [
            ["VIX volatility index", "Decision-month daily average", "Implied uncertainty and risk appetite"],
            ["Baa corporate yield minus 10Y Treasury", "Decision-month average spread", "Credit conditions and default compensation"],
            ["10Y minus 2Y yield", "Decision-month average spread", "Expected business conditions and term compensation"],
            ["2Y yield change", "Change in monthly average yield", "Monetary conditions and discount-rate adjustment"],
            ["CPI inflation", "Monthly change of preceding released level", "Price pressure and policy response"],
            ["Unemployment", "Preceding released monthly level", "Labour-market slack and demand conditions"],
            ["IP growth", "Monthly change of preceding released level", "Real activity and expected cash-flow conditions"],
        ], [1.55, 2.40, 2.55], "Sources: supplied FRED series and the release calendar. 2Y and 10Y denote two- and ten-year Treasury maturities. Percentage-point transformations are used except the VIX index and unemployment level. Year-on-year IP and CPI define states, not replacements for continuous monthly changes.", ["results/states/model_calendar.csv", "data"])
    if kind == "specification":
        return ("Table 3 Documented specification choices", ["Choice", "Retained setting", "Relation to observed outcomes"], [
            ["Earlier allocation", "Raw positive forecasts across six factors", "Already examined; retained as a sensitivity"],
            ["Other mappings", "Variance scaling, fixed market exposure, rank weights and volatility ceiling", "Explored with outcomes available; list is not a complete search log"],
            ["Current allocation", "Five long–short factors; 36-month variances; 50% shrinkage; 40% cap", "Selected after outcome inspection; frozen for the reported comparisons"],
            ["State rules", "Fixed and expanding year-on-year medians", "Retrospective choices; neither restores Expansion in the original evaluation"],
            ["Target window", "175 months including the former 87-month validation period", "Previously examined outcomes; not an independent holdout"],
            ["Penalties", "Initial training time-series cross-validation", "Later validation and evaluation errors excluded from penalty selection"],
        ], [1.35, 2.35, 2.80], "Exact decision dates and a complete preregistered search history are unavailable. The table records documented choices and does not claim that the listed mappings exhaust all experiments.", ["specification.json"])
    if kind == "forecast_accuracy":
        rows=[]
        for factor in FACTORS:
            f=select(data["forecasts"], Design=CURRENT, Factor=factor)
            rows.append([factor.replace("MKT_RF", "MKT−RF")] + [format_r2(model_r2(f,col)) for col in ["Macro_OLS", "Macro_Ridge", "Macro_Elastic_Net", "Macro"]])
        return ("Table 4 Forecast accuracy against the allocation historical mean", ["Factor", "OLS R² %", "Ridge R² %", "Elastic Net R² %", "Average R² %"], rows, [1.10, 1.25, 1.25, 1.55, 1.35], "175 targets including three mean fallbacks. Every column uses the same allocation historical mean; positive R² means lower MSE. Matched-label OLS tests appear separately.", ["results/us/forecasts.csv"])
    if kind == "performance":
        rows=[]
        for strategy in STRATEGIES:
            r=perf[strategy]
            rows.append([LABEL[strategy], number(r["Annualized_Mean"],2,100),number(r["Annualized_Volatility"],2,100),number(r["Sharpe"]),number(r["Synthetic_CAGR"],2,100),number(r["Max_Drawdown"],2,100)])
        return ("Table 5 Primary allocation over 175 target months", ["Strategy", "Mean %", "Volatility %", "Sharpe", "CAGR %", "Max DD %"], rows, [1.75,.95,1.05,.80,.95,1.0], "January 2012–July 2026; five long–short factors; current-vintage inputs; 10-basis-point allocation charge. Mean and volatility are annualized. CAGR and drawdown use synthetic compounded excess returns and omit funding and factor-internal trading.", ["results/us/metrics.csv"])
    if kind == "uncertainty":
        rows=[]
        for cohort in ["Extended_175", "Original_88"]:
            for comparison in ["Macro_minus_Historical_Mean", "Macro_M1_YoY_minus_Macro", "Macro_M1_ExpandingYoY_minus_Macro"]:
                filters=dict(Design=CURRENT,Cohort=cohort,Allocation=PRIMARY,Sample="Full",Comparison=comparison,Cost_bps="10",HAC_Lags="3")
                r=unique(data["intervals"], Block_Length="6", **filters)
                m=unique(data["mde"], **filters)
                t=unique(data["mean_tests"], **filters)
                rows.append([LABEL[comparison],r["N_Months"],number(r["Sharpe_Difference"]),f'[{number(r["CI_Low"])}, {number(r["CI_High"])}]',number(m["Approx_MDE_Annual_Sharpe_80pct"]),number(t["Mean"],3,12),number(t["Holm_p_Within_Displayed_Comparison_Family"],3)])
        return ("Table 6 Paired portfolio differences and uncertainty", ["Comparison", "N", "Δ Sharpe", "95% interval", "MDE", "Mean gain pp", "Holm p"], rows, [1.85,.35,.60,1.35,.55,.85,.60], "Sharpe intervals: 3,000 paired moving-block percentile-t draws, six-month blocks, HAC3 studentization, net 10-basis-point allocation charges. MDE: approximate annual Sharpe difference detectable at two-sided 5% and 80% power. Mean gains are annual percentage points; Holm adjusts three mean comparisons within each window. No correction covers the wider specification search.", ["results/us/studentized_intervals.csv", "results/us/mde.csv", "results/us/mean_difference_hac.csv"])
    if kind == "states":
        rows=[]
        for variant, label in [("FrozenYoY","Fixed"),("ExpandingYoY","Expanding")]:
            for split in ["Training","Validation","Evaluation"]:
                r=select(data["states"],Design=CURRENT,Variant=variant,Split=split)
                first=r[0]
                counts=[unique(r,Regime=x)["N"] for x in ["Expansion","Overheating","Slowdown","Stagflation"]]
                rows.append([f'{label} / {split}',*counts,str(int(first["N_All_Signals"])-int(first["N_Valid_States"])),first["N_All_Signals"]])
        return ("Table 7 State coverage on signal dates", ["Rule and window", "Expansion", "Overheating", "Slowdown", "Stagflation", "Missing", "Total"], rows, [1.70,.90,1.05,.95,.95,.65,.60], "Training: March 1990–November 2011. Former validation: December 2011–February 2019. Original evaluation: March 2019–June 2026. Each return target ends one month later. Counts use unique calendar signals, not factor rows.", ["results/states/state_frequencies.csv"])
    if kind == "robustness":
        rows=[]
        for cost in [0,25,50]:
            common=dict(Design=CURRENT,Cohort="Extended_175",Allocation=PRIMARY,Sample="Full",Cost_bps=str(cost))
            hm=unique(data["metrics"],Strategy="Historical_Mean",**common)
            macro=unique(data["metrics"],Strategy="Macro",**common)
            rows.append([f'{cost} bp allocation charge',number(hm["Sharpe"]),number(macro["Sharpe"]),number(float(macro["Sharpe"])-float(hm["Sharpe"])),number((float(macro["Annualized_Mean"])-float(hm["Annualized_Mean"]))*100)])
        common=dict(Design=CURRENT,Cohort="Extended_175",Allocation=PRIMARY,Sample="Full",Allocation_Cost_bps="10",MOM_Internal_Cost_bps_per_month="65.0")
        hm=unique(data["mom"],Strategy="Historical_Mean",**common);macro=unique(data["mom"],Strategy="Macro",**common)
        rows.append(["10 bp plus MOM 65 bp",number(hm["Sharpe"]),number(macro["Sharpe"]),number(float(macro["Sharpe"])-float(hm["Sharpe"])),number(float(macro["Annual_Mean_pp"])-float(hm["Annual_Mean_pp"]))])
        return ("Table 8 Cost sensitivity over 175 months", ["Scenario", "Historical mean Sharpe", "Macro Sharpe", "Δ Sharpe", "Mean gain pp"], rows,[2.1,1.30,1.1,.85,1.05], "Allocation charges are basis points per half absolute target-weight change. The momentum scenario adds a hypothetical 65-basis-point monthly charge per unit MOM weight to each strategy’s own weights. Mean gain is Macro minus historical mean in annual percentage points. IP input and penalty comparisons appear in Appendix C.", ["results/us/metrics.csv", "results/diagnostics/mom_internal_cost_sensitivity.csv"])
    if kind == "hypothesis_outcomes":
        h1_min={c:min(float(r["Holm_p_Six_Factors"]) for r in select(data["h1"], Design=CURRENT, Cohort=c, Estimator="OLS", Method="Clark_West", HAC_Lags="3")) for c in ["Extended_175","Original_88"]}
        h2=lambda c,s:unique(data["h2"], Design=CURRENT, Cohort=c, State_Model=s, Factor="RMW", Estimator="OLS", Method="Clark_West", HAC_Lags="3")
        ci=unique(data["intervals"], Design=CURRENT, Cohort="Extended_175", Allocation=PRIMARY, Cost_bps="10", Sample="Full", Comparison="Macro_minus_Historical_Mean", Block_Length="6", HAC_Lags="3")
        rows=[
            ["Do continuous signals improve forecasts? (H1-F)","OLS versus a mean fitted on the same past labels","No broad improvement after six-factor correction"],
            ["Do continuous signals improve allocation? (H1-A)",f'Macro versus the identically mapped historical mean; ΔSharpe {number(ci["Sharpe_Difference"])}',"Small positive estimate; 95% interval spans zero"],
            ["Do states improve forecasts? (H2-F)","State OLS versus continuous OLS on identical labels","RMW conditional rejection in 175 months; none in 88"],
            ["Do states improve allocation? (H2-A)","State portfolios versus continuous Macro","No positive gain; original fixed-state comparison deteriorates"],
        ]
        rows.append(["Does the allocation pattern extend beyond the US?", "Frozen allocation rule on UK and German Ridge forecasts", "Small unresolved UK increment; near-zero negative German estimate; state portfolios lower"])
        return ("Table 1 Research questions and evidence",["Question","Comparison","Finding"],rows,[2.25,2.2,2.05],"Decisions use a 5% level. US forecast tests use one-sided matched OLS Clark–West, HAC3 and separate six-factor Holm families; the 175/88 target windows have 172/85 eligible origins. Portfolio intervals use paired six-month blocks and HAC3. International allocations use country-specific Ridge forecasts and their matched historical means; their mean-test Holm family contains four country–contrast comparisons per HAC lag.",["results/us/direct_h1_forecast_loss.csv","results/us/direct_h2_forecast_loss.csv","results/us/studentized_intervals.csv","results/cross_market/allocation_summary.csv"])
    if kind == "mapping_family":
        rows=[]
        for mapping in data["selection_summary"]["mappings"]:
            r=unique(data["mapping_comparisons"], Cohort="Extended_175", Mapping=mapping)
            t=unique(data["mapping_turnover"],Cohort="Extended_175",Mapping=mapping)
            rows.append([MAPPING_LABEL[mapping],number(r["Historical_Mean_Sharpe"]),number(r["Macro_Sharpe"]),number(r["Sharpe_Difference"]),number(r["Annual_Mean_Gain_pp"]),number(t["Macro_Mean_Monthly_One_Way_Turnover"],1,100)])
        macro=[unique(data["reality_check"], Cohort="Extended_175",Family="Macro_10_Mappings",Block_Length=str(b)) for b in [3,6,12]]
        combined=[unique(data["reality_check"], Cohort="Extended_175",Family="Information_Models_30",Block_Length=str(b)) for b in [3,6,12]]
        macro_p=" / ".join(number(r["Reality_Check_p"]) for r in macro)
        combined_p=" / ".join(number(r["Reality_Check_p"]) for r in combined)
        note=("175 targets, current-vintage forecasts, net 10 bp allocation charges. Variance and rank mappings use 50% equal-weight shrinkage and a 40% cap; raw mappings normalize positive forecasts. A 5% annual-volatility ceiling scales risky weights downward, with residual zero-excess-return cash. Fixed-market mappings reserve 1/6 for MKT−RF. Mean gain is Macro minus the identically mapped historical mean in annual pp. "
              f'Reality Check p-values for 3 / 6 / 12-month blocks: Macro ten-mapping family {macro_p}; thirty Macro/state combinations {combined_p}. '
              f'Paired non-circular moving blocks, {macro[1]["Bootstrap_Draws"]} draws, unstudentized maximum of centered mean gains. The primary mapping’s maximum-adjusted p is {number(macro[1]["Primary_Max_Adjusted_p"])} at six-month blocks. The correction covers these documented reconstructed mappings, conditional on fitted paths.')
        note+=" Turnover is Macro’s mean monthly one-way target-weight change, in %, with free initial entry. Mappings can change scaling and constraints jointly."
        return ("Table D1 All documented allocation mappings",["Mapping","Mean SR","Macro SR","Δ Sharpe","Mean gain pp","Turnover %"],rows,[2.25,.80,.80,.80,.90,.95],note,["results/selection/mapping_comparisons.csv","results/selection/reality_check.csv","results/robustness/mapping_turnover.csv"])
    if kind == "complexity":
        rows=[]
        for factor in FACTORS:
            ridge=unique(data["complexity"],Design=CURRENT,Factor=factor,Model="Ridge")
            en=unique(data["complexity"],Design=CURRENT,Factor=factor,Model="Elastic_Net")
            alpha=lambda v:"10⁶" if float(v)==1e6 else number(v,2)
            rows.append([factor.replace("MKT_RF","MKT−RF"),alpha(ridge["Alpha"]),number(ridge["Slope_Effective_DF"],4),number(en["Alpha"],4),str(int(float(en["Active_Predictor_Count"])))])
        return ("Table D2 Initial continuous-model complexity",["Factor","Ridge penalty","Ridge slope df","EN penalty","EN active slopes"],rows,[.95,1.20,1.30,1.40,1.65],"261 initial current-vintage training signals. Ridge slope df is trace[Z(Z′Z + αI)⁻¹Z′], for the centered standardized design and fixed penalty; add one for the unpenalized intercept. EN active coefficients use a 10⁻¹⁰ threshold and are a complexity proxy, not an exact effective-df estimate. EN mixing ratios and all refit coefficients remain in the numerical outputs.",["results/diagnostics/initial_effective_df.csv"])
    if kind == "covid_mappings":
        rows=[]
        for mapping in [PRIMARY,"Raw6_Sensitivity"]:
            r=unique(data["covid_mappings"],Design=CURRENT,Cohort="Extended_175",Allocation=mapping,Cost_bps="10",Period="2020_Mar_Dec")
            rows.append([MAPPING_LABEL[mapping],r["N"],number(r["Macro_Arithmetic_Sum_pp"],2),number(r["HM_Arithmetic_Sum_pp"],2),number(r["Macro_minus_HM_Arithmetic_pp"],2)])
        return ("Table D3 Same COVID episode under two allocation rules",["Mapping","N","Macro sum pp","Mean sum pp","Difference pp"],rows,[2.05,.45,1.25,1.25,1.50],"March–December 2020, current-vintage inputs, 10 bp allocation charges. These are sums of monthly arithmetic returns and paired differences. They are distinct from separately compounded portfolio returns or a difference in compounded wealth.",["results/interpretation/event_2020_mapping_comparison.csv"])
    if kind == "international":
        rows=[]
        for country in ["UK","DE"]:
            for theme in ["size","value","investment","profitability","momentum"]:
                m=unique(data["intl"],country=country,theme=theme,model="Macro")
                s=unique(data["intl"],country=country,theme=theme,model="Macro+Regime")
                rows.append([f'{country} {theme}',format_r2(m["oos_r2"]),format_r2(s["oos_r2"]),format_rms_bp(m["prediction_minus_mean_rms_bp"]),format_rms_bp(s["prediction_minus_mean_rms_bp"])])
        return ("Table B1 UK and German Ridge forecasts",["Country and theme","Macro R² %","State R² %","Macro RMS bp","State RMS bp"],rows,[1.80,1.1,1.1,1.2,1.2],"UK 133 and Germany 122 targets; IP reference lag two, CLIFS lag two, CPI lag one. R² compares local historical means. RMS is the prediction’s departure from its matched training mean, not its return forecast error. Signed R² values smaller than 0.01% are shown as +<0.01 or −<0.01.",["results/international/model_shrinkage_diagnostics.csv"])
    if kind == "ip_comparison":
        rows=[]
        designs=[(CURRENT,"Current IP / common penalties"),("Partial_realtime_IP_Current_Penalties","Historical IP / common penalties"),("Partial_realtime_IP","Historical IP / retuned penalties")]
        for cohort,n in [("Extended_175","175"),("Original_88","88")]:
            for design,label in designs:
                t=unique(data["ip_mean_tests"],Design=design,Cohort=cohort,HAC_Lags="3")
                ci=unique(data["ip_intervals"],Design=design,Cohort=cohort)
                rows.append([n,label,number(t["Annual_Mean_Gain_pp"]),number(ci["Sharpe_Difference"]),f'[{number(ci["CI_Low"])}, {number(ci["CI_High"])}]'])
        return ("Table C1 IP input and penalty comparisons",["N","Input and estimation setting","Mean gain pp","Δ Sharpe","95% interval"],rows,[.40,2.65,.95,.80,1.70],"Macro minus its identically mapped historical mean, primary five-factor allocation, 10 bp sleeve charges. Common penalties are the current-IP initial-training choices, applied unchanged to both input calendars. The historical-IP retuned design uses its own initial-training choices. Sharpe intervals use 3,000 paired six-month moving-block draws and HAC3 studentization. Each contrast conditions on its generated paths.",["results/robustness/ip_design_mean_hac.csv","results/robustness/ip_design_sharpe_intervals.csv","results/robustness/ip_parameter_audit.csv"])
    if kind == "ip_mean_sensitivity":
        rows=[]
        for design,label in [(CURRENT,"Current IP"),("Partial_realtime_IP_Current_Penalties","Historical IP / common"),("Partial_realtime_IP","Historical IP / retuned")]:
            for lag in [3,6,12]:
                r=unique(data["ip_mean_tests"],Design=design,Cohort="Extended_175",HAC_Lags=str(lag))
                archived=r["Archived_Within_Design_Three_Comparison_Holm_p"]
                rows.append([label,str(lag),number(r["p_Two_Sided"],4),number(archived,4) if archived else "—",number(r["Holm_p_Three_Design_Contrasts"],4)])
        return ("Table C2 Mean-return dependence sensitivities over 175 months",["IP setting","HAC lag","Raw p","Portfolio-family Holm p","IP-family Holm p"],rows,[2.30,.70,.90,1.30,1.30],"Two-sided tests of zero expected monthly Macro-minus-mean return. Portfolio-family Holm adjusts three contrasts within each existing design and lag: Macro minus mean, fixed state minus Macro, expanding state minus Macro. IP-family Holm instead adjusts the three Macro-minus-mean input/penalty settings within each cohort and lag. The common-penalty IP supplement has no state fits, so its portfolio-family entry is undefined. Neither correction covers selection across HAC lags or the complete specification search. HAC3 is primary; HAC6 and HAC12 are sensitivities. Corresponding 88-month values are in the mapped result file.",["results/robustness/ip_design_mean_hac.csv"])
    if kind == "training_slopes":
        rows=[]
        predictors=list(dict.fromkeys(r["Predictor"] for r in data["training_slopes"]))
        labels={"VIX":"VIX","Credit_Spread":"Credit spread","Term_Spread":"Term spread","Treasury_2Y_change":"2Y yield change","Inflation_lag":"CPI inflation","Unemployment_lag":"Unemployment","IP_Growth_lag":"IP growth"}
        for predictor in predictors:
            cells=[]
            for factor in FACTORS:
                r=unique(data["training_slopes"],Factor=factor,Predictor=predictor)
                marker="†" if float(r["Holm_p_42_coefficients"])<.05 else "*" if float(r["HAC_p"])<.05 else ""
                cells.append(number(r["Coefficient"],3)+marker)
            rows.append([labels.get(predictor,predictor.replace("_"," ")),*cells])
        return ("Table E1 Initial-training OLS slope associations",["Predictor","MKT−RF","SMB","HML","RMW","CMA","MOM"],rows,[1.55,.825,.825,.825,.825,.825,.825],"261 initial training signals, next-month factor return in percentage points. Coefficients use raw predictor units: one VIX point, one unemployment percentage point, or one percentage point of the specified growth, inflation, spread or yield change. * denotes two-sided HAC3 p<0.05 before correction; † denotes Holm p<0.05 across all 42 slopes. Four raw slopes and one corrected slope pass 5%; the surviving association is SMB with credit spread. These are initial-sample conditional associations, not causal channels or out-of-sample skill. Cross-validation selects penalties using forecast loss, not these p-values.",["results/robustness/training_slopes.csv"])
    if kind == "rmw_state_means":
        rows=[]
        splits=["Training_261","Former_validation_87","Original_88","Extended_175"]
        for variant,label in [("FixedYoY","Fixed"),("ExpandingYoY","Expanding")]:
            for state in ["Expansion","Overheating","Slowdown","Stagflation"]:
                cells=[]
                for split in splits:
                    r=unique(data["rmw_state_means"],Variant=variant,Split=split,Regime=state)
                    cells.append(r["N_RMW_Outcomes"]+" / "+(number(r["RMW_Next_Month_Mean_pct"],3) if r["RMW_Next_Month_Mean_pct"] else "—"))
                rows.append([label,state,*cells])
        return ("Table E2 Observed next-month RMW returns by state",["Rule","State","Training N / mean %","Former 87 N / mean %","Original 88 N / mean %","Extended 175 N / mean %"],rows,[.70,1.15,1.15,1.15,1.15,1.20],"Cells give the number of state-labeled targets and the mean subsequent monthly operating-profitability factor return in %. Training targets end April 1990–December 2011; former validation January 2012–March 2019; original evaluation April 2019–July 2026. Two state labels are missing in each later evaluation window; state-labeled months may include a continuous-model fallback, so the original 86 labeled targets differ from 85 matched forecast origins. An absent state has N=0 and an undefined mean (—). Means are descriptive conditional associations; small unequal cells do not identify a defensive or causal mechanism. Dispersion and continuous-model eligibility counts remain in the mapped CSV.",["results/robustness/rmw_state_means.csv"])
    if kind == "glossary":
        return ("Table A1 Terms used in this report",["Term","Meaning in the experiment"],[
            ["Signal and target", "Month-end information and the return ending in the following month"],
            ["Historical mean", "Average of completed past factor returns; matched-label means are identified separately"],
            ["Factor weight", "Share of a synthetic long–short factor return; not a long-only stock holding"],
            ["Fallback", "Historical-mean forecast used when a required current predictor is unavailable"],
            ["HAC", "Covariance estimate allowing heteroskedasticity and dependence across specified calendar lags"],
            ["HAC3 / HAC6 / HAC12", "HAC covariance uses 3, 6 or 12 calendar-month lags; these are not bootstrap block lengths"],
            ["Clark–West test", "One-sided adjusted squared-loss comparison for a nested forecast model; penalized and ensemble applications are exploratory approximations"],
            ["Sleeve", "An allocation to one synthetic factor or theme return; sleeve turnover excludes trading within its long and short stock portfolios"],
            ["Holm adjustment", "Sequential multiplicity correction within an explicitly stated comparison family"],
            ["Studentized interval", "Bootstrap interval formed from an estimate divided by a recomputed standard error"],
            ["Percentile-t", "Quantiles of bootstrap estimate departures divided by their recomputed standard errors"],
            ["MDE", "Local approximate difference detectable at a specified significance level and power"],
            ["Reality Check", "A joint bootstrap comparison of the largest matched mean-return gain in a stated candidate family"],
            ["Spanning regression", "Macro strategy returns regressed on historical-mean strategy returns and an intercept"],
        ],[1.6,4.9],"Returns and differences are labelled explicitly: % for return levels, pp for percentage-point differences, and bp for basis points. One percentage point equals 100 basis points. Sharpe ratios and their differences are unitless annualized summaries.",[])
    raise ValueError(kind)


NS_M="http://schemas.openxmlformats.org/officeDocument/2006/math"


def m_el(name, **attrs):
    x=OxmlElement("m:"+name)
    for k,v in attrs.items(): x.set(qn("m:"+k),str(v))
    return x


def mr(text, normal=False):
    r=m_el("r")
    if normal:
        prop=m_el("rPr");prop.append(m_el("nor"));r.append(prop)
    prop=OxmlElement("w:rPr"); fonts=OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"),"Cambria Math");fonts.set(qn("w:hAnsi"),"Cambria Math");prop.append(fonts)
    size=OxmlElement("w:sz");size.set(qn("w:val"),"22");prop.append(size);r.append(prop)
    t=m_el("t");t.text=str(text)
    if t.text != t.text.strip():
        t.set("{http://www.w3.org/XML/1998/namespace}space","preserve")
    r.append(t);return r


def mg(*items):
    out=[]
    for i in items:
        if isinstance(i,(str,int,float)):out.append(mr(i))
        elif isinstance(i,(list,tuple)):out.extend(mg(*i))
        else:out.append(i)
    return out


def mc(name,*items):
    x=m_el(name);x.extend(mg(*items));return x


def sub(x,i):return mc("sSub",mc("e",x),mc("sub",i))
def sup(x,i):return mc("sSup",mc("e",x),mc("sup",i))
def frac(x,y):return mc("f",mc("num",x),mc("den",y))
def sqrt(x):
    prop=m_el("radPr");prop.append(m_el("degHide",val="1"))
    return mc("rad",prop,m_el("deg"),mc("e",x))
def hat(x):
    prop=m_el("accPr");prop.append(m_el("chr",val="̂"))
    return mc("acc",prop,mc("e",x))
def bar(x):
    prop=m_el("barPr");prop.append(m_el("pos",val="top"))
    return mc("bar",prop,mc("e",x))
def delim(x,left="(",right=")"):
    prop=m_el("dPr");prop.extend([m_el("begChr",val=left),m_el("endChr",val=right),m_el("grow",val="1")])
    return mc("d",prop,mc("e",x))


def equations(kind):
    pred=lambda model:sub(hat("μ"),["j,t+1|t",",",mr(model,True)])
    if kind=="forecast":
        return [mg(sub(hat("μ"),"j,t+1|t"),"=",frac([pred("OLS"),"+",pred("Ridge"),"+",pred("EN")],"3"))]
    if kind=="weights":
        return [
            mg(sub("q","j,t"),"=",mr("max",True),delim(["0,",frac(sub(hat("μ"),"j,t+1|t"),sup(sub(hat("σ"),"j,t,36"),"2"))]),",  ",sub("p","j,t"),"=",frac(sub("q","j,t"),[sub("∑","k"),sub("q","k,t")]),mr("  if ",True),sub("∑","k"),sub("q","k,t"),">0"),
            mg(sub("w","t"),"=",sub(mr("C",True),"0.40"),delim(["0.5",sub("p","t"),"+",frac("0.5","5"),mr("1",True)])),
        ]
    if kind=="return":
        return [mg(sub("r",[mr("net",True),",t+1"]),"=∑",sub("w","j,t"),sub("r","j,t+1"),"−",frac("c","10,000"),"×",frac("1","2"),"∑",delim([sub("w","j,t"),"−",sub("w","j,t−1")],"|","|"))]
    if kind=="statistics":
        return [mg(mr("SR",True),"=",sqrt("12"),"×",frac(bar("r"),sub("s","r"))), mg(mr("MDE",True),"≈",delim([sub("z","0.975"),"+",sub("z","0.80")]),sub(mr("SE",True),"ΔSR"),"≈2.8016",sub(mr("SE",True),"ΔSR"))]
    raise ValueError(kind)


def set_font(run, size=11, bold=False, italic=False):
    run.font.name="Times New Roman";run.font.size=Pt(size);run.bold=bold;run.italic=italic
    run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();fonts=rpr.find(qn("w:rFonts"))
    if fonts is None:fonts=OxmlElement("w:rFonts");rpr.append(fonts)
    fonts.set(qn("w:ascii"),"Times New Roman");fonts.set(qn("w:hAnsi"),"Times New Roman")


def configure(doc):
    # Remove decorative paragraph borders inherited from the default template.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.80);sec.right_margin=Inches(.80)
    sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
    for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Caption"]:
        s=doc.styles[name];s.font.name="Times New Roman";s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
        s.element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman");s.element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman")
    normal=doc.styles["Normal"];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.widow_control=True
    title=doc.styles["Title"];title.font.size=Pt(20);title.font.bold=True;title.paragraph_format.space_after=Pt(6)
    sub=doc.styles["Subtitle"];sub.font.size=Pt(11);sub.paragraph_format.space_after=Pt(6)
    h1=doc.styles["Heading 1"];h1.font.size=Pt(14);h1.font.bold=True;h1.paragraph_format.space_before=Pt(13);h1.paragraph_format.space_after=Pt(6)
    h2=doc.styles["Heading 2"];h2.font.size=Pt(11.5);h2.font.bold=True;h2.paragraph_format.space_before=Pt(9);h2.paragraph_format.space_after=Pt(4)
    caption=doc.styles["Caption"];caption.font.size=Pt(10);caption.font.italic=False;caption.paragraph_format.space_before=Pt(8);caption.paragraph_format.space_after=Pt(5)
    caption.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("UTSG 7394   |   "),9)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    core=doc.core_properties
    core.author="";core.last_modified_by="";core.comments="";core.subject="US equity factor prediction and allocation with macro-financial signals"
    core.created=datetime(2026,10,8,tzinfo=timezone.utc);core.modified=core.created


def write_paragraph(doc,text,style=None):
    p=doc.add_paragraph(style=style)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(text)
    return p


def render_reference(doc, text):
    """Format APA author/year text with journal title and volume italics."""
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
    p.paragraph_format.line_spacing=1.02;p.paragraph_format.space_after=Pt(5)
    # Locate the publication details from the right, then the immediately
    # preceding journal name. Article titles ending in ? or ! do not receive
    # an extra full stop in APA, so matching only '. ' can absorb the title.
    publication_matches=list(re.finditer(r", (\d+)(\(\d+\))?, ([\d–−-]+)\.",text))
    publication=publication_matches[-1] if publication_matches else None
    journal=re.search(r"[.!?]\s+([^.!?\n]+)$",text[:publication.start()]) if publication else None
    website=re.search(r"\(n\.d\.(?:-[a-z])?\)\. (.+?)\. Retrieved ",text)
    italic_ranges=[]
    if journal:
        italic_ranges=[(journal.start(1),journal.end(1)),(publication.start(1),publication.end(1))]
    elif website:
        italic_ranges=[(website.start(1),website.end(1))]
    cursor=0
    for lo,hi in italic_ranges:
        set_font(p.add_run(text[cursor:lo]),9.5)
        set_font(p.add_run(text[lo:hi]),9.5,italic=True)
        cursor=hi
    set_font(p.add_run(text[cursor:]),9.5)
    return p


def render_table(doc, title, columns, rows, widths, note):
    write_paragraph(doc,title,"Caption")
    table=doc.add_table(rows=1,cols=len(columns));table.autofit=False
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for i, text in enumerate(columns):table.rows[0].cells[i].text=text
    for row in rows:
        cells=table.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text);cells[i].width=Inches(widths[i])
    for ri,row in enumerate(table.rows):
        pr=row._tr.get_or_add_trPr()
        if ri==0:pr.append(OxmlElement("w:tblHeader"))
        # Keeping each short row intact avoids ambiguous continuation.
        pr.append(OxmlElement("w:cantSplit"))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement("w:tcMar")
            for side,value in [("top",70),("bottom",70),("left",90),("right",90)]:
                x=OxmlElement("w:"+side);x.set(qn("w:w"),str(value));x.set(qn("w:type"),"dxa");margins.append(x)
            tcpr.append(margins)
            borders=OxmlElement("w:tcBorders")
            for side in ["top","left","bottom","right","insideH","insideV"]:
                border=OxmlElement("w:"+side);border.set(qn("w:val"),"single");border.set(qn("w:sz"),"4");border.set(qn("w:color"),"D9D9D9");borders.append(border)
            tcpr.append(borders)
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"E8EDF2" if ri==0 else "FFFFFF");tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.02;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(columns)==3 or len(columns)==2 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:set_font(run,9.5,bold=ri==0)
    p=write_paragraph(doc,note)
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.03
    for run in p.runs:set_font(run,9)
    if len(rows) <= 8:
        # Keep short tables with their definitions and avoid unnecessary continuations.
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next=True


def replace_values(text,values):
    def repl(match):return values[match.group(1)]
    return re.sub(r"\{\{(\w+)\}\}",repl,text)


def source_formatted_document(base):
    """Restore the source Word's formatting, without any earlier report body."""
    template = json.loads((base / "report/format_template.json").read_text(encoding="utf-8"))
    memory = BytesIO()
    with ZipFile(memory, "w") as package:
        for name, encoded in template["parts"].items():
            package.writestr(name, base64.b64decode(encoded))
    memory.seek(0)
    doc = Document(memory)
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    core = doc.core_properties
    core.author = ""; core.last_modified_by = ""; core.comments = ""
    core.subject = "US equity factor allocation with exploratory UK and German checks"
    core.modified = datetime(2026, 10, 8, tzinfo=timezone.utc)
    return doc


def link_cross_references(doc):
    """Add Word navigation anchors and reject references with no target."""
    targets = {}
    for p in doc.paragraphs:
        match = re.match(r"^(Appendix [A-F]|Table (?:[A-F])?\d+)\b", p.text)
        if match and p.style.name in {"Caption", "Heading 1", "Heading 2"}:
            label = match.group(1)
            if label in targets:
                raise ValueError("Duplicate report cross-reference target: " + label)
            name = label.replace(" ", "_")
            identifier = str(len(targets) + 1)
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), identifier); start.set(qn("w:name"), name)
            end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), identifier)
            position = 1 if p._p.pPr is not None else 0
            p._p.insert(position, start); p._p.append(end)
            targets[label] = name
    references = []
    first_appendix_mentions = []
    in_bibliography = False
    for p in doc.paragraphs:
        if p.text == "References": in_bibliography = True
        if p.text == "Appendices": in_bibliography = False
        if in_bibliography: continue
        if p.style.name in {"Caption", "Heading 1", "Heading 2", "Title", "Subtitle"}:
            continue
        text = p.text
        matches = list(re.finditer(r"\b(Appendix [A-F]|Table (?:[A-F])?\d+)\b", text))
        if not matches:
            continue
        if any(p._p.iter(qn("m:oMath"))):
            raise ValueError("Cross-reference unexpected inside native equation")
        props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        for child in list(p._p):
            if child.tag != qn("w:pPr"):
                p._p.remove(child)
        cursor = 0
        for match in matches:
            label = match.group(1)
            if label not in targets:
                raise ValueError("Unresolved report cross-reference: " + label)
            if match.start() > cursor:
                run = p.add_run(text[cursor:match.start()])
                if props is not None: run._r.insert(0, deepcopy(props))
            link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), targets[label])
            run = OxmlElement("w:r")
            if props is not None: run.append(deepcopy(props))
            value = OxmlElement("w:t"); value.text = label; run.append(value)
            link.append(run); p._p.append(link)
            references.append(label)
            if label.startswith("Appendix") and label not in first_appendix_mentions:
                first_appendix_mentions.append(label)
            cursor = match.end()
        if cursor < len(text):
            run = p.add_run(text[cursor:])
            if props is not None: run._r.insert(0, deepcopy(props))
    expected = ["Appendix " + letter for letter in "ABCDE"]
    if first_appendix_mentions[:5] != expected:
        raise ValueError("Appendices do not follow first substantive mention: " + str(first_appendix_mentions))
    return {"all_passed": True, "targets": targets, "linked_references": references,
            "first_appendix_mentions": first_appendix_mentions,
            "scope": "Body paragraphs and table notes; AI declaration F need not be cited"}


def build(base: Path):
    base=Path(base).resolve()
    manuscript=json.loads((base/"report/manuscript.json").read_text(encoding="utf-8"))
    data,perf,values=data_context(base)
    main_words=sum(len(re.findall(r"\b[\w’−]+\b",replace_values(x["text"],values))) for s in manuscript["sections"] for x in s["blocks"] if x["type"]=="p")
    doc=source_formatted_document(base);doc.core_properties.title=manuscript["title"]
    title_display=manuscript["title"].replace("Improve US Equity", "Improve\nUS Equity")
    write_paragraph(doc,title_display,"Title")
    write_paragraph(doc,manuscript["subtitle"],"Subtitle")
    if manuscript.get("course"):
        p=write_paragraph(doc,manuscript["course"])
        for r in p.runs:set_font(r,10)
    write_paragraph(doc,"Abstract","Heading 2")
    write_paragraph(doc,replace_values(manuscript["abstract"],values))
    equation_number=0;source_map=[]
    def block(item):
        nonlocal equation_number
        kind=item["type"]
        if kind in ["p","h2"]:
            write_paragraph(doc,replace_values(item["text"],values),"Heading 2" if kind=="h2" else None)
        elif kind=="table":
            title,cols,rows,widths,note,sources=table_data(item["id"],data,perf)
            render_table(doc,title,cols,rows,widths,note)
            modules={"predictors":"src/us_design.py; src/states.py", "specification":"specification.json", "forecast_accuracy":"src/us_analysis.py", "performance":"src/us_analysis.py", "uncertainty":"src/us_analysis.py", "states":"src/states.py", "robustness":"src/us_analysis.py; src/diagnostics.py", "hypothesis_outcomes":"src/us_analysis.py; src/states.py", "mapping_family":"src/selection.py; src/robustness.py", "complexity":"src/diagnostics.py", "covid_mappings":"src/interpretation.py", "international":"src/international_analysis.py", "glossary":"src/build_report.py", "ip_comparison":"src/robustness.py", "ip_mean_sensitivity":"src/robustness.py", "training_slopes":"src/training.py; src/robustness.py", "rmw_state_means":"src/states.py; src/robustness.py", "cross_market_allocation":"src/cross_market.py", "international_sources":"documentation/data_source_decisions.csv", "international_allocation_uncertainty":"src/cross_market.py"}
            source_map.append({"item":title,"builder":"src/build_report.py","analysis":modules[item["id"]],"inputs":sources})
        elif kind=="equation":
            for nodes in equations(item["id"]):
                equation_number+=1
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(8)
                p.paragraph_format.keep_together=True
                para=m_el("oMathPara");prop=m_el("oMathParaPr");prop.append(m_el("jc",val="center"));para.append(prop)
                formula=mc("oMath",nodes,mr(f"   ({equation_number})",True));para.append(formula);p._p.append(para)
        elif kind=="figure":
            id_=item["id"]
            captions={
                "figure6":"Figure 1 Cumulative raw squared-error gains for six factor forecasts. Solid lines compare continuous OLS with its matched-label mean across 172 eligible targets; dashed lines compare the actual three-model average with the allocation mean across all 175 targets, including three fallbacks. Positive values indicate lower squared error. The different benchmarks and target sets are identified explicitly; these lines contain no Clark–West adjustment and are not significance tests.",
                "figure1":"Figure 2 Primary synthetic excess-return paths and cumulative Macro-minus-historical-mean arithmetic difference. All strategies use 10 bp allocation charges. The dashed boundary begins the original April 2019 target window. The lower panel sums paired monthly arithmetic differences and is distinct from compounded wealth differences.",
                "figure2":"Figure 3 Macro weights in the five long–short factors. Dates mark return targets; weights use information at the preceding decision month. The market-factor weight remains zero.",
                "figure3":"Figure 4 Growth–inflation states on signal dates under fixed and expanding medians. Boundaries separate initial training, former validation and original evaluation. Missing observations remain visible.",
                "figure7":"Figure 5 Annual absolute and incremental returns under the primary allocation and 10 bp charges. The upper panel separately compounds Macro and historical-mean returns; the lower panel sums paired monthly Macro-minus-mean differences in pp. The lower bars add to the full arithmetic increment. The 2026 bars cover January–July only. Absolute return concentration in 2022 differs from the annual pattern of relative gains.",
            }
            path=base/f"results/figures/{id_}.png"
            if not path.is_file():raise FileNotFoundError(path)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
            p.add_run().add_picture(str(path),width=Inches(6.55))
            c=write_paragraph(doc,captions[id_]);c.paragraph_format.space_after=Pt(8);c.paragraph_format.line_spacing=1.04
            for r in c.runs:set_font(r,9)
            figure_inputs={"figure1":"results/us/returns.csv", "figure2":"results/us/weights.csv", "figure3":"results/states/model_calendar.csv", "figure6":"results/diagnostics/cumulative_forecast_loss.csv", "figure7":"results/robustness/annual_incremental_returns.csv"}
            numbers={"figure6":1,"figure1":2,"figure2":3,"figure3":4,"figure7":5}
            analysis="src/robustness.py" if id_=="figure7" else "src/diagnostics.py" if id_=="figure6" else "src/states.py" if id_=="figure3" else "src/us_analysis.py"
            source_map.append({"item":f'Figure {numbers[id_]} ({id_}.png)',"builder":"src/plot_figures.py","analysis":analysis,"inputs":[figure_inputs[id_]]})
        else:raise ValueError(kind)
    for section in manuscript["sections"]:
        write_paragraph(doc,section["title"],"Heading 1")
        for item in section["blocks"]:block(item)
    references_heading=write_paragraph(doc,"References","Heading 1")
    references_heading.paragraph_format.page_break_before=True
    if not manuscript["references"]:raise ValueError("Verified cited references have not been supplied")
    for ref in manuscript["references"]:
        render_reference(doc,ref)
    write_paragraph(doc,"Appendices","Heading 1")
    for item in manuscript["appendix"]:block(item)
    cross_references=link_cross_references(doc)
    output=base/"report/Final_Report.docx";doc.save(output)
    (base/"results/report_cross_reference_audit.json").write_text(json.dumps(cross_references,indent=2))
    overview={
        "report":"report/Final_Report.docx","main_prose_words":main_words,
        "main_word_count_scope":"Paragraph prose only; excludes headings, abstract, tables, figures, references and appendices",
        "abstract_words":len(replace_values(manuscript["abstract"],values).split()),
        "main_tables":sum(x["type"]=="table" for s in manuscript["sections"] for x in s["blocks"]),
        "appendix_tables":sum(x["type"]=="table" for x in manuscript["appendix"]),
        "figures":sum(x["type"]=="figure" for s in manuscript["sections"] for x in s["blocks"]),"native_equations":equation_number,
        "references":len(manuscript["references"]),"numerical_text":values,
    }
    (base/"results/report_build_summary.json").write_text(json.dumps(overview,indent=2,ensure_ascii=False),encoding="utf-8")
    (base/"report/report_sources.json").write_text(json.dumps(source_map,indent=2,ensure_ascii=False),encoding="utf-8")
    documentation=base/"documentation";documentation.mkdir(exist_ok=True)
    with (documentation/"report_output_map.csv").open("w",encoding="utf-8-sig",newline="") as handle:
        writer=csv.DictWriter(handle,fieldnames=["Report_item","Source_data_or_result","Analysis_code","Presentation_code"])
        writer.writeheader()
        for row in source_map:
            writer.writerow({"Report_item":row["item"],"Source_data_or_result":"; ".join(row["inputs"]),"Analysis_code":row["analysis"],"Presentation_code":row["builder"]})
    print(f"Built report/Final_Report.docx: {overview['main_prose_words']} main prose words, {overview['main_tables']} main tables, {overview['appendix_tables']} appendix tables, {overview['figures']} figures, {equation_number} editable equations")
    return output


def run_analysis(base: Path):
    return build(base)


if __name__=="__main__":
    build(Path(__file__).resolve().parents[1])
